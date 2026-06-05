import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation
import io

def extract_pdf_text(file_bytes: bytes) -> str:
    """Extracts text from a PDF file using PyMuPDF."""
    text = ""
    with fitz.open(stream=file_bytes, filetype="pdf") as doc:
        for page in doc:
            text += page.get_text()
    return text

def extract_docx_text(file_bytes: bytes) -> str:
    """Extracts text from a DOCX file using python-docx."""
    doc = Document(io.BytesIO(file_bytes))
    return "\n".join([para.text for para in doc.paragraphs])

def extract_pptx_text(file_bytes: bytes) -> str:
    """Extracts text from a PPTX file using python-pptx."""
    prs = Presentation(io.BytesIO(file_bytes))
    text_runs = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text_frame") and shape.text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        text_runs.append(run.text)
    return "\n".join(text_runs)
